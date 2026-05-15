"""
Mise à jour de la base de données Excel papers_db.xlsx

Ajoute pour chaque schéma trouvé :
  - Numéro de figure (ex: "1.2", "4.5")
  - Code BibTeX complet
  
Gère la numérotation par papier (1.1, 1.2, 1.3 puis 2.1, 2.2, etc.)

Dépendance : openpyxl  →  pip install openpyxl
"""

from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

from utils import build_bibtex, logger


class ExcelDatabaseUpdater:
    """Ajoute les figures et BibTeX à la base de données Excel."""

    def __init__(self, excel_path: str | Path):
        self.excel_path = Path(excel_path)
        if not self.excel_path.exists():
            raise FileNotFoundError(f"Excel non trouvé : {excel_path}")
        self.wb = load_workbook(self.excel_path)
        self.ws = self.wb.active

    # ─────────────────────────────────────────────────────────────────
    #  PUBLIC
    # ─────────────────────────────────────────────────────────────────

    def add_diagrams(self, diagrams: list[dict]) -> None:
        """
        Ajoute tous les schémas à la base.
        Gère la numérotation par papier.
        """
        # Regrouper par papier (par DOI/titre)
        diagrams_by_paper = self._group_by_paper(diagrams)

        # Trouver la dernière ligne et le dernier numéro de papier
        last_row = self._find_last_row()
        last_paper_num = self._find_last_paper_number()
        current_paper_num = last_paper_num + 1

        for paper_id, paper_diagrams in diagrams_by_paper.items():
            for fig_idx, diagram in enumerate(paper_diagrams, 1):
                last_row += 1
                self._insert_row(
                    row_num=last_row,
                    paper_number=current_paper_num,
                    figure_number=fig_idx,
                    diagram=diagram,
                )
            current_paper_num += 1

        self.wb.save(str(self.excel_path))
        logger.info(f"✓ Excel mis à jour : {self.excel_path}")
        logger.info(f"  {sum(len(d) for d in diagrams_by_paper.values())} schémas insérés")

    # ─────────────────────────────────────────────────────────────────
    #  PRIVATE
    # ─────────────────────────────────────────────────────────────────

    def _group_by_paper(self, diagrams: list[dict]) -> dict[str, list[dict]]:
        """Groupe les schémas par papier unique."""
        grouped = {}
        for diagram in diagrams:
            paper = diagram["paper"]
            # Clé unique : DOI ou titre normalisé
            key = (
                paper.get("doi") or
                paper.get("title", "").lower()[:50]
            )
            if key not in grouped:
                grouped[key] = []
            grouped[key].append(diagram)
        return grouped

    def _find_last_row(self) -> int:
        """Trouve le numéro de la dernière ligne remplie."""
        for row in range(self.ws.max_row, 0, -1):
            if self.ws.cell(row, 1).value is not None:
                return row
        return 0

    def _find_last_paper_number(self) -> int:
        """Récupère le plus grand numéro de papier (colonne A)."""
        max_num = 0
        for row in range(2, self.ws.max_row + 1):  # Sauter le header
            val = self.ws.cell(row, 1).value
            if val and isinstance(val, int):
                max_num = max(max_num, val)
        return max_num

    def _insert_row(
        self,
        row_num: int,
        paper_number: int,
        figure_number: int,
        diagram: dict,
    ) -> None:
        """Insère une ligne pour un schéma."""
        paper = diagram["paper"]

        # Colonne A : Numéro de papier
        self.ws.cell(row_num, 1).value = paper_number

        # Colonne B : Numéro de figure (ex: "4.2")
        figure_num_str = f"{paper_number}.{figure_number}"
        self.ws.cell(row_num, 2).value = figure_num_str

        # Colonne C : BibTeX complet
        bibtex = build_bibtex(paper)
        self.ws.cell(row_num, 3).value = bibtex

        # Formater les cellules
        self._format_row(row_num)

    def _format_row(self, row_num: int) -> None:
        """Applique le formatage à la nouvelle ligne."""
        # Colonne A : nombre centré
        cell_a = self.ws.cell(row_num, 1)
        cell_a.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)

        # Colonne B : numéro centré
        cell_b = self.ws.cell(row_num, 2)
        cell_b.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)

        # Colonne C : texte wrappé
        cell_c = self.ws.cell(row_num, 3)
        cell_c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell_c.font = Font(name="Arial", size=10)

        # Auto-adjust column width (si possible)
        # Note : openpyxl ne supporte pas bien l'auto-width, on fixe à la main
        self.ws.column_dimensions["C"].width = 80


def update_excel_from_diagrams(
    excel_path: str | Path,
    diagrams: list[dict],
) -> None:
    """Wrapper public pour mettre à jour Excel."""
    updater = ExcelDatabaseUpdater(excel_path)
    updater.add_diagrams(diagrams)
